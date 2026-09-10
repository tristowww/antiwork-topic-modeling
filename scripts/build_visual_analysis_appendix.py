"""Build a shareable visual appendix from aggregate topic-modeling evidence."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "outputs" / "v2_clean_min25" / "figures"
SENTIMENT_FIGURES = ROOT / "outputs" / "sentiment" / "twitter_roberta_base_sentiment_latest"
OUT = ROOT / "deliverables" / "Antiwork_Visual_Analysis_Appendix.docx"
INK = "000000"
MUTED = "4B5563"

FIGURE_PAGES = [
    ("From collection to estimates", "brief_collection_to_results_flow.png", 6.4, "Collection, filtering, topic discovery, review, and descriptive estimation. The appendix uses aggregate evidence only; raw post text remains local."),
    ("Topic assignment coverage", "monthly_assignment_coverage.png", 6.35, "The share of filtered posts assigned to a non-outlier BERTopic cluster by month. Unassigned posts are a coverage limitation rather than a substantive category."),
    ("Main theme trends", "brief_key_theme_monthly_trends.png", 6.35, "Three-month rolling prevalence for the themes central to the descriptive interpretation. Each series is divided by all filtered posts in that month."),
    ("Ten largest reviewed subthemes by month", "top10_reviewed_subtheme_monthly_prevalence.png", 6.35, "Exact month-by-month prevalence for the ten largest included BERTopic clusters, which are subthemes nested within the nine reviewed themes. Panel n values are full-period cluster sizes."),
    ("All reviewed theme trends", "candidate_theme_monthly_prevalence.png", 6.35, "Monthly prevalence for all included candidate themes. The codebook covers 18,506 posts, or 19.0% of the filtered corpus."),
    ("Theme prevalence by rolling window", "candidate_theme_rolling_window_prevalence.png", 5.6, "Exploratory theme prevalence across the four March-to-February rolling windows. These are shares of all filtered posts, not the full content taxonomy."),
    ("Y1 and Y4 theme comparison", "candidate_theme_y1_y4_comparison.png", 6.35, "Exploratory Y1-to-Y4 comparison of candidate-theme prevalence. Y1 is labeled as the Great Resignation period (March 2021-February 2022); connecting lines show descriptive percentage-point change, not causal effects."),
    ("Anchor topic prevalence", "monthly_anchor_topic_prevalence.png", 6.35, "Monthly prevalence of selected anchor topics from the full-corpus model. This provides topic-level context alongside the reviewed-theme map."),
    ("Largest topic prevalence", "monthly_topic_prevalence_top20_full.png", 6.35, "Monthly prevalence across the 20 most prevalent full-corpus model topics. This is exploratory descriptive context, not a codebook of final themes."),
    ("Theme and subtheme hierarchy", "reviewed_theme_hierarchy.png", 5.1, "The nine reviewed candidate themes, their included subthemes, and full-corpus cluster sizes. Node area scales with cluster n."),
    ("Whole-corpus polarity by rolling window", SENTIMENT_FIGURES / "sentiment_by_rolling_window.png", 6.35, "Context-aware whole-post RoBERTa polarity estimates for all 97,254 retained posts. They are model-based, not human-validated or management-targeted sentiment labels."),
    ("Polarity within reviewed themes", SENTIMENT_FIGURES / "theme_analysis" / "candidate_theme_negative_probability_by_month.png", 6.35, "Three-month rolling mean of model-estimated negative whole-post polarity within each reviewed theme. It is conditional on theme assignment and distinct from theme prevalence."),
    ("Window alignment Y1 to Y2", "alignment_Y1_2021-22_vs_Y2_2022-23.png", 4.7, "Conservative top-term alignment between adjacent rolling-window topic fits. It is supporting context for stability, not evidence of a persistent construct."),
    ("Window alignment Y2 to Y3", "alignment_Y2_2022-23_vs_Y3_2023-24.png", 4.7, "Conservative top-term alignment between adjacent rolling-window topic fits. It is supporting context for stability, not evidence of a persistent construct."),
    ("Window alignment Y3 to Y4", "alignment_Y3_2023-24_vs_Y4_2024-25.png", 4.7, "Conservative top-term alignment between adjacent rolling-window topic fits. It is supporting context for stability, not evidence of a persistent construct."),
    ("Window alignment Y1 to Y4", "alignment_Y1_2021-22_vs_Y4_2024-25.png", 4.7, "A long-range alignment comparison between the first and fourth rolling-window fits. Use this as a robustness diagnostic, not a claim of stable themes across all years."),
]


def set_run(run, *, size: float, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def paragraph(doc, text: str, *, size: float = 10.5, color: str = INK, bold: bool = False, italic: bool = False, after: float = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("Antiwork topic modeling visual appendix"), size=8.2, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("Page "), size=8.2, color=MUTED)
    add_page_number(footer)
    for name, size in (("Normal", 10.5), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)


def build() -> None:
    def resolve_figure(path_or_name: Path | str) -> Path:
        return path_or_name if isinstance(path_or_name, Path) else FIGURES / path_or_name

    missing = [str(resolve_figure(path_or_name)) for _, path_or_name, _, _ in FIGURE_PAGES if not resolve_figure(path_or_name).is_file()]
    if missing:
        raise FileNotFoundError(f"Missing figures: {missing}")
    doc = Document()
    configure(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(title.add_run("Antiwork Visual Analysis Appendix"), size=24, bold=True)
    paragraph(doc, "Exploratory topic-modeling figures for collaborator review", size=12, color=MUTED, after=16)
    paragraph(doc, "This appendix gathers the available aggregate graphics used to assess the corpus, document the reviewed theme map, describe longitudinal patterns, and evaluate topic-model stability. It is intended to accompany the analysis brief and manuscript, not replace their methods and limitations.", size=11.2, after=10)
    paragraph(doc, "Scope note. The primary corpus contains 97,254 management-term r/antiwork posts from March 2021 through February 2025. Topic prevalence figures use all filtered posts as the denominator unless their captions state otherwise. The reviewed candidate-theme codebook contains 27 clusters and covers 18,506 posts (19.0% of the filtered corpus).", size=10.5, after=10)
    paragraph(doc, "The appendix also includes the completed fixed RoBERTa whole-post polarity check. These estimates provide descriptive evidence about the problem-oriented sampling frame; they are not human-validated, management-targeted, or causal measures. Per-post labels and the four-post smoke-test graphic remain local.", size=10.5, italic=True, color=MUTED, after=0)

    for number, (heading, path_or_name, width, caption) in enumerate(FIGURE_PAGES, start=1):
        doc.add_page_break()
        h = doc.add_paragraph(style="Heading 1")
        set_run(h.add_run(f"Figure {number}  {heading}"), size=15, bold=True)
        image = doc.add_picture(str(resolve_figure(path_or_name)), width=Inches(width))
        image.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph(doc, caption, size=9.2, color=MUTED, italic=True, after=0)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
