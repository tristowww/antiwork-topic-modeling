"""Build the collaborator-ready exploratory analysis brief.

The brief only uses aggregate, shareable evidence. Raw post text, sampled
excerpts, embeddings, and local-only analysis artifacts are intentionally absent.
"""
from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DOCX = OUT_DIR / "Antiwork_Exploratory_Analysis_Brief.docx"
REPO_URL = "https://github.com/tristowww/antiwork-topic-modeling"
CODEBOOK_PATH = ROOT / "planning" / "CANDIDATE_THEME_CODEBOOK.csv"
TOPIC_ASSIGNMENTS_PATH = ROOT / "outputs" / "v2_clean_min25" / "bertopic_full_document_topics.csv"
WINDOW_PREVALENCE_PATH = ROOT / "outputs" / "v2_clean_min25" / "analysis" / "candidate_theme_window_prevalence.csv"
RAW_ARCHIVE_SUBMISSION_RECORDS = 577_190
SENTIMENT_DIR = ROOT / "outputs" / "sentiment" / "twitter_roberta_base_sentiment_latest"
SENTIMENT_SUMMARY_PATH = SENTIMENT_DIR / "sentiment_summary.json"

NAVY = "17365D"
BLUE = "2E74B5"
MID_BLUE = "4978A8"
MUTED = "5D6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
INK = "1F2937"


def set_run(run, size: float = 11, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    element = tc_pr.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        tc_pr.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in {"top": 80, "start": 120, "bottom": 80, "end": 120}.items():
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table._tbl.tblPr.first_child_found_in("w:tblW").set(qn("w:w"), str(sum(widths)))
    table._tbl.tblPr.first_child_found_in("w:tblW").set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def paragraph(doc, text: str = "", *, size: float = 10.5, color: str = INK, bold: bool = False, italic: bool = False, after: float = 6, style: str = "Normal"):
    p = doc.add_paragraph(style=style)
    if text:
        set_run(p.add_run(text), size, color, bold, italic)
    p.paragraph_format.space_after = Pt(after)
    return p


def heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    set_run(p.add_run(text), 15 if level == 1 else 12.5, NAVY, True)


def bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    set_run(p.add_run(text), 10.2)


def hyperlink(paragraph_obj, text: str, url: str) -> None:
    relation_id = paragraph_obj.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph_obj._p.append(link)


def callout_pair(doc, left_title: str, left_text: str, right_title: str, right_text: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_widths(table, [4680, 4680])
    for cell, title, body in zip(table.rows[0].cells, (left_title, right_title), (left_text, right_text)):
        shade(cell, "FAFBFC")
        set_run(cell.paragraphs[0].add_run(title), 9.6, NAVY, True)
        p = cell.add_paragraph()
        set_run(p.add_run(body), 9.2)


def corpus_metrics() -> dict[str, int]:
    with TOPIC_ASSIGNMENTS_PATH.open(newline="", encoding="utf-8") as handle:
        assignments = list(csv.DictReader(handle))
    with CODEBOOK_PATH.open(newline="", encoding="utf-8") as handle:
        included_topics = {int(row["topic"]) for row in csv.DictReader(handle) if row["decision"] == "include"}
    return {
        "filtered_posts": len(assignments),
        "assigned_posts": sum(int(row["topic"]) >= 0 for row in assignments),
        "mapped_posts": sum(int(row["topic"]) in included_topics for row in assignments),
    }


def metric_strip(doc) -> None:
    metrics = corpus_metrics()
    table = doc.add_table(rows=1, cols=4)
    set_table_widths(table, [2340, 2340, 2340, 2340])
    for cell, (value, label) in zip(table.rows[0].cells, [
        (f"{metrics['filtered_posts']:,}", "filtered posts"),
        ("199", "non-outlier clusters"),
        (f"{metrics['assigned_posts'] / metrics['filtered_posts']:.1%}", f"topic assignment\nn={metrics['assigned_posts']:,} of {metrics['filtered_posts']:,}"),
        (f"{metrics['mapped_posts'] / metrics['filtered_posts']:.1%}", f"theme-map coverage\nn={metrics['mapped_posts']:,} of {metrics['filtered_posts']:,}"),
    ]):
        shade(cell, "FAFBFC")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cell.paragraphs[0].add_run(value), 13, NAVY, True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(label), 8.3, MUTED)


def findings_table(doc) -> None:
    table = doc.add_table(rows=1, cols=4)
    set_table_widths(table, [3300, 1500, 1500, 3060])
    with WINDOW_PREVALENCE_PATH.open(newline="", encoding="utf-8") as handle:
        prevalence = list(csv.DictReader(handle))
    lookup = {(row["candidate_theme"], row["rolling_window"]): row for row in prevalence}
    y1_n = int(next(row["total_posts"] for row in prevalence if row["rolling_window"] == "Y1_2021-22"))
    y4_n = int(next(row["total_posts"] for row in prevalence if row["rolling_window"] == "Y4_2024-25"))
    for cell, label in zip(table.rows[0].cells, ["Theme", f"Y1\n(n={y1_n:,})", f"Y4\n(n={y4_n:,})", "Interpretation"]):
        shade(cell, LIGHT_BLUE)
        set_run(cell.paragraphs[0].add_run(label), 9, NAVY, True)
    interpretation = {
        "Health safety and attendance": "Early COVID/illness concentration contracts.",
        "Scheduling hours and time off": "Largest mapped theme in Y4.",
        "Compensation and pay rights": "Remains substantial in every window after a Y2 high.",
        "Career precarity and exit": "Broadly stable across four rolling windows.",
    }
    rows = []
    for theme, note in interpretation.items():
        y1 = lookup[(theme, "Y1_2021-22")]
        y4 = lookup[(theme, "Y4_2024-25")]
        rows.append((theme, f"{float(y1['share_all_posts']):.2%}\n(n={int(y1['post_count']):,})", f"{float(y4['share_all_posts']):.2%}\n(n={int(y4['post_count']):,})", note))
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        if index % 2:
            for cell in cells:
                shade(cell, "FAFBFC")
        for cell, value in zip(cells, row):
            set_run(cell.paragraphs[0].add_run(value), 8.5)


def rolling_window_counts_table(doc) -> None:
    """Show the numerator behind every rolling-window prevalence percentage."""
    with WINDOW_PREVALENCE_PATH.open(newline="", encoding="utf-8") as handle:
        prevalence = list(csv.DictReader(handle))
    order = ["Y1_2021-22", "Y2_2022-23", "Y3_2023-24", "Y4_2024-25"]
    totals = {window: int(next(row["total_posts"] for row in prevalence if row["rolling_window"] == window)) for window in order}
    theme_order = (
        Counter({theme: sum(int(row["post_count"]) for row in prevalence if row["candidate_theme"] == theme) for theme in {row["candidate_theme"] for row in prevalence}})
        .most_common()
    )
    lookup = {(row["candidate_theme"], row["rolling_window"]): row for row in prevalence}
    table = doc.add_table(rows=1, cols=5)
    set_table_widths(table, [3000, 1590, 1590, 1590, 1590])
    headers = ["Candidate theme", *[f"{window.replace('_', ' ')}\n(n={totals[window]:,})" for window in order]]
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_run(cell.paragraphs[0].add_run(label), 7.8, NAVY, True)
    for index, (theme, _) in enumerate(theme_order):
        cells = table.add_row().cells
        if index % 2:
            for cell in cells:
                shade(cell, "FAFBFC")
        set_run(cells[0].paragraphs[0].add_run(theme), 7.5, INK, True)
        for cell, window in zip(cells[1:], order):
            row = lookup[(theme, window)]
            value = f"{float(row['share_all_posts']):.2%}\n(n={int(row['post_count']):,})"
            set_run(cell.paragraphs[0].add_run(value), 7.5, INK)


def robustness_table(doc) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_widths(table, [2500, 6860])
    for cell, label in zip(table.rows[0].cells, ["Check", "Result"]):
        shade(cell, LIGHT_BLUE)
        set_run(cell.paragraphs[0].add_run(label), 9, NAVY, True)
    rows = [
        ("Corpus audit", "Inputs contain 97,254 matching unique post IDs; the management-term filter re-matches 100% of retained posts."),
        ("Seed stability", "Across three UMAP seeds on the same 10,000-post stratified sample, joint-inlier ARI was 0.900 to 0.966."),
        ("Window alignment", "Conservative one-to-one term matching is supporting context, not proof of substantive persistence."),
        ("Text sensitivity", "Titles-only clustering was stable but materially different from title-plus-body; title-plus-body remains the primary representation."),
    ]
    for index, (label, result) in enumerate(rows):
        cells = table.add_row().cells
        if index % 2:
            for cell in cells:
                shade(cell, "FAFBFC")
        set_run(cells[0].paragraphs[0].add_run(label), 8.5, NAVY, True)
        set_run(cells[1].paragraphs[0].add_run(result), 8.5)


def reviewed_cluster_table(doc, rows: list[dict[str, str]]) -> None:
    """Render a compact, inspectable appendix table for the reviewed clusters."""
    table = doc.add_table(rows=1, cols=5)
    set_table_widths(table, [620, 620, 2050, 4500, 1570])
    headers = ["ID", "n", "Representative terms", "Candidate theme and subtheme", "Decision"]
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        set_run(cell.paragraphs[0].add_run(label), 8.0, NAVY, True)

    for index, row in enumerate(rows):
        cells = table.add_row().cells
        if index % 2:
            for cell in cells:
                shade(cell, "FAFBFC")
        terms = " ".join(row["top_terms"].split()[:6])
        candidate = f"{row['candidate_theme']}\n{row['candidate_subtheme']}"
        decision = f"{row['decision'].title()}\n{row['confidence']} confidence"
        if row["decision"] == "include":
            shade(cells[3], "E2F0D9")
        else:
            shade(cells[3], LIGHT_GRAY)
        for cell, value in zip(cells, (row["topic"], f"{int(row['cluster_n']):,}", terms, candidate, decision)):
            set_run(cell.paragraphs[0].add_run(value), 7.5, INK, cell in cells[:2])


def load_reviewed_clusters() -> list[dict[str, str]]:
    with CODEBOOK_PATH.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    with TOPIC_ASSIGNMENTS_PATH.open(newline="", encoding="utf-8") as handle:
        counts = Counter(int(row["topic"]) for row in csv.DictReader(handle))
    for row in rows:
        row["cluster_n"] = str(counts[int(row["topic"])])
    return sorted(rows, key=lambda row: int(row["topic"]))


def load_sentiment_summary() -> dict[str, object]:
    return json.loads(SENTIMENT_SUMMARY_PATH.read_text(encoding="utf-8"))


def sentiment_outputs_available() -> bool:
    return all(path.is_file() for path in [
        SENTIMENT_SUMMARY_PATH,
        SENTIMENT_DIR / "sentiment_by_rolling_window.png",
        SENTIMENT_DIR / "theme_analysis" / "candidate_theme_negative_probability_by_month.png",
    ])


def configure(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.top_margin = Inches(0.76)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    for header in (section.header, section.even_page_header):
        h = header.paragraphs[0]
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run(h.add_run("Antiwork topic modeling | Analysis brief"), 8.2, MUTED, False)
    for footer in (section.footer, section.even_page_footer):
        f = footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(f.add_run("Updated 11 September 2026 | Exploratory evidence package"), 8.2, MUTED)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level, size, before, after, color in ((1, 15, 12, 6, NAVY), (2, 12.5, 9, 4, NAVY), (3, 11.5, 7, 3, NAVY)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure(doc)
    paragraph(doc, "Analysis brief", size=9.2, color=MUTED, after=3)
    paragraph(doc, "Antiwork Topic Modeling", size=22, color=NAVY, bold=True, after=2)
    paragraph(doc, "Exploratory longitudinal analysis for collaborator review", size=11.5, color=MUTED, after=12)
    metric_strip(doc)
    has_sentiment = sentiment_outputs_available()
    heading(doc, "Study summary")
    paragraph(doc, f"Across {RAW_ARCHIVE_SUBMISSION_RECORDS:,} archived r/antiwork submission records retrieved from March 2021 through February 2025, the management-term filter retained 97,254 unique posts for analysis. The updated BERTopic analysis identifies a descriptive shift away from COVID-centered health and safety discussion and toward scheduling and time-off discussion. Compensation, career precarity, and recruitment remain present across the four rolling windows.")
    callout_pair(doc, "Completed analysis", "Rebuilt the analysis with BERTopic; exported document-level assignments and four rolling-window fits; finalized a 27-cluster, nine-theme codebook; and produced figures and reproducibility materials.", "Interpretive scope", "The figures show the share of all filtered posts captured by included candidate-theme clusters. The 27-cluster map is not a full taxonomy or a complete estimate of management-related content.")
    heading(doc, "Selected findings")
    findings_table(doc)
    paragraph(doc, "All percentages divide posts in included candidate-theme clusters by every filtered post in the same rolling window. They show the portion of the filtered corpus captured by this 27-cluster map, not a complete measure of all management-related content.", size=8.5, color=MUTED, italic=True, after=4)
    heading(doc, "Data collection", 2)
    paragraph(doc, f"Public r/antiwork submissions were retrieved through the Arctic Shift archive API in 48 monthly UTC batches from 1 March 2021 through 1 March 2025. This collection yielded {RAW_ARCHIVE_SUBMISSION_RECORDS:,} archived submission records before filtering. The analysis retained 97,254 unique posts whose title or body matched the established management terms boss, manager, supervisor, or team lead, including plural variants. The resulting corpus is an archive-based sample of public discourse, not a representative sample of employees or an author-level panel. Archive records may differ from content displayed on Reddit after collection.", size=9.1, after=0)
    doc.add_page_break()

    heading(doc, "Analytic workflow")
    paragraph(doc, "The analysis proceeded as follows. Corpus checks, model-based topic discovery, descriptive prevalence estimates, and robustness evidence are reported separately.", size=9.5, color=MUTED, after=8)
    steps = [
        ("Collect, filter, and audit the corpus", f"Retrieved {RAW_ARCHIVE_SUBMISSION_RECORDS:,} archived submission records across 48 monthly pulls, then applied the management-term filter and confirmed 97,254 unique retained post IDs. A duplicate audit found 1.2% residual duplicate non-placeholder texts; this is reported rather than silently removed."),
        ("Construct the primary text field", "Used post title plus body after handling deleted and placeholder content. A title-only run was retained as a sensitivity analysis, not substituted for the primary representation."),
        ("Discover topics with a current topic-modeling pipeline", "Embedded documents with all-MiniLM-L6-v2, reduced the embedding space with UMAP, clustered with HDBSCAN, and represented topics with BERTopic c-TF-IDF. The full-corpus model produced 199 non-outlier clusters and assigned 38.6% of posts."),
        ("Document the theme map", "Reviewed the largest 30 clusters and retained 27 interpretable clusters in a nine-theme candidate codebook. Generic, deleted, and community-meta clusters were excluded, yielding 18,506 mapped posts (19.0% of the full filtered corpus)."),
        ("Estimate descriptive longitudinal patterns", "Calculated each theme's share of every filtered post in four March-to-February rolling windows. The figures in this brief show the main monthly patterns, the full reviewed-theme comparison, and monthly assignment coverage."),
        ("Test stability and state the boundary", "Ran seed-stability checks, title-only sensitivity, and one-to-one top-term alignment checks across rolling-window fits. A fixed RoBERTa whole-post polarity check then scored all 97,254 retained posts and was joined to the reviewed map without refitting either model; causal, human-validated sentiment, and full-conversation prevalence claims remain out of scope."),
    ]
    for number, (title, detail) in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        set_run(p.add_run(f"{number}. {title}. "), 10.0, NAVY, True)
        set_run(p.add_run(detail), 9.5, INK)
    paragraph(doc, "The private repository includes the analysis scripts, codebook, diagnostics, and aggregate evidence used in this brief. Raw posts and text-level samples remain local.", size=8.6, color=MUTED, italic=True, after=0)
    heading(doc, "From collection to estimates", 2)
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "brief_collection_to_results_flow.png"), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 1. The collection, filtering, modeling, review, and estimation sequence used for the exploratory analysis.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Monthly patterns in the main themes")
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "brief_key_theme_monthly_trends.png"), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 2. Three-month rolling prevalence of the themes that anchor the main interpretation. Shares use all filtered posts each month as the denominator.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Month-by-month longitudinal patterns in the ten largest reviewed clusters")
    paragraph(doc, "The top-level candidate map contains nine themes, so a literal top-ten theme chart would be mislabeled. For the requested month-by-month longitudinal view, the figure below shows the ten largest included BERTopic clusters, nested within the nine-theme map and labeled as subthemes rather than as additional themes. Each panel uses the full filtered monthly corpus as its denominator.", size=9.1, after=4)
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "top10_reviewed_subtheme_monthly_prevalence.png"), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 3. Month-by-month longitudinal prevalence of the ten largest reviewed clusters, with each cluster's parent candidate theme in its label. The shaded period is the Y1 Great Resignation window, March 2021 through February 2022. Panel n values are full-period cluster sizes.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "All reviewed themes by rolling window")
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "candidate_theme_rolling_window_prevalence.png"), width=Inches(4.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 4. Exploratory prevalence by rolling window. The legend gives each window denominator; Table 1 gives the numerator n behind every displayed percentage. The 27-cluster map excludes generic, deleted, and community-meta clusters.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Counts behind rolling-window percentages")
    paragraph(doc, "Each cell gives the theme share of all filtered posts in that rolling window followed by its numerator n. The column-header n is the all-filtered-post denominator for that window.", size=9.1, after=5)
    rolling_window_counts_table(doc)
    paragraph(doc, "Table 1. Counts and prevalence for every reviewed candidate theme across the four rolling windows. Values describe the documented 27-cluster map, not all management-related discussion.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Reviewed theme hierarchy")
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "reviewed_theme_hierarchy.png"), width=Inches(5.05))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 5. The nine candidate themes, their included subthemes, and full-corpus cluster sizes. Node area scales with n; this documented map covers 18,506 posts, or 19.0% of the filtered corpus.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Exploratory Y1 and Y4 comparison")
    paragraph(doc, "The comparison below contrasts the user-defined Great Resignation period, March 2021 through February 2022, with the most recent March 2024 through February 2025 window. It is a descriptive comparison of the reviewed theme map, not an estimate of the Great Resignation's effect.", size=9.1, after=4)
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "candidate_theme_y1_y4_comparison.png"), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 6. Candidate-theme prevalence in Y1 and Y4. The legend gives the all-filtered-post denominator for each window; Table 1 gives each theme numerator n; labels show descriptive percentage-point differences.", size=8.2, color=MUTED, italic=True, after=0)
    doc.add_page_break()

    heading(doc, "Whole-corpus polarity check")
    if has_sentiment:
        sentiment = load_sentiment_summary()
        hard_shares = sentiment["hard_label_shares"]
        posterior = sentiment["mean_posterior_probability"]
        callout_pair(
            doc,
            "Context-aware model",
            "A fixed social-media RoBERTa classifier estimated whole-post negative, neutral, and positive polarity across every retained post. This is a contextual transformer-based estimate, rather than a lexicon method such as VADER.",
            "Interpretive boundary",
            "The check is model-based, not human-validated or targeted specifically to management. It tests whether the intentionally problem-oriented sample is predominantly negative under this classifier; it does not estimate average employee sentiment.",
        )
        paragraph(
            doc,
            f"Across {sentiment['scored_posts']:,} scored posts, the model's hard labels were {hard_shares['negative']:.1%} negative, {hard_shares['neutral']:.1%} neutral, and {hard_shares['positive']:.1%} positive. Mean posterior probabilities were {posterior['negative']:.1%}, {posterior['neutral']:.1%}, and {posterior['positive']:.1%}, respectively. Mean maximum class probability was {sentiment['mean_max_probability']:.1%}; {sentiment['truncated_share']:.1%} of posts reached the 512-token limit.",
            size=9.1,
            after=4,
        )
        paragraph(doc, "The hard negative-label share was 59.8% (n=16,923 of 28,276) in Y1, the Great Resignation window, and 67.7% (n=8,096 of 11,960) in Y4. This descriptive increase is not a causal estimate and does not establish a change in employee sentiment.", size=9.1, after=4)
        doc.add_picture(str(SENTIMENT_DIR / "sentiment_by_rolling_window.png"), width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph(doc, "Figure 7. Whole-corpus RoBERTa polarity estimates across the four rolling windows. The hard-label and mean-posterior views are shown together because model uncertainty is material to interpretation.", size=8.2, color=MUTED, italic=True, after=0)
        doc.add_page_break()

        heading(doc, "Polarity within reviewed themes")
        paragraph(doc, "The figure below joins the fixed post-level RoBERTa estimates to the existing reviewed BERTopic map. It does not refit either model. Values describe whole-post negative-polarity probability among posts assigned to each theme, while the prevalence figures above retain all filtered posts as their denominator.", size=9.1, after=4)
        doc.add_picture(str(SENTIMENT_DIR / "theme_analysis" / "candidate_theme_negative_probability_by_month.png"), width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph(doc, "Figure 8. Three-month rolling mean of modeled negative-polarity probability within the nine reviewed themes. The theme map covers 18,506 posts, or 19.0% of the filtered corpus.", size=8.2, color=MUTED, italic=True, after=0)
    else:
        paragraph(doc, "The planned descriptive polarity check uses a fixed social-media RoBERTa classifier rather than a lexicon method. Its full-corpus run and quality checks are still incomplete, so no sentiment result or graphic is included in this brief.", size=9.1, after=0)
    doc.add_page_break()

    heading(doc, "Analytic approach and scope")
    callout_pair(doc, "Primary analysis", "Sentence-transformer embeddings (all-MiniLM-L6-v2), UMAP, HDBSCAN, BERTopic topic representations, and a full-corpus model for longitudinal prevalence. Per-window fits serve as stability diagnostics.", "Coverage of estimates", "The model assigns 38.6% of posts to non-outlier clusters. The final theme map covers 18,506 posts: 19.0% of all filtered posts and 49.3% of assigned posts.")
    heading(doc, "Assignment coverage", 2)
    doc.add_picture(str(ROOT / "outputs" / "v2_clean_min25" / "figures" / "monthly_assignment_coverage.png"), width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "Figure 9. Assignment coverage varied from 33.1% to 44.0% by month. Outliers are a coverage limitation, not a substantive category.", size=8.2, color=MUTED, italic=True)
    doc.add_page_break()
    heading(doc, "Robustness checks", 2)
    robustness_table(doc)
    paragraph(doc, "A larger all-mpnet-base-v2 encoder check was attempted but did not complete within the available CPU bound, so it is deferred and not treated as evidence.", size=8.5, color=MUTED, italic=True)

    heading(doc, "Interpreting the results", 2)
    callout_pair(doc, "Observed patterns", "Within the mapped themes, health, safety, and attendance discussion declined over the four windows, while scheduling, hours, and time-off discussion increased. Compensation, career precarity, and recruitment remained visible throughout.", "What the results support", "The topic patterns describe changes in the composition of management-related public discussion and identify domains for follow-up with representative employee data. They do not establish employee sentiment, turnover drivers, or policy effects.")

    heading(doc, "Remaining work")
    callout_pair(doc, "Available for drafting", "Results, figures, a theme codebook, coverage diagnostics, seed stability, text-mode sensitivity, whole-corpus polarity estimates, and a defined interpretive scope are ready for the manuscript.", "Future validation", "The planned human-validation study below would test construct validity beyond the current exploratory evidence. It has not been conducted and is not a basis for any current result.")
    heading(doc, "Planned hybrid human-validation study", 2)
    paragraph(doc, "The strongest next step is a two-stage hybrid validation study, not a single preset-topic yes-or-no check. The design separates discovery of human-meaningful concepts from direct evaluation of the frozen BERTopic theme map, while keeping coders blinded to model assignments during primary coding.", size=9.1, after=4)
    paragraph(doc, "1. Blinded inductive coding. Draw a stratified exploratory subsample across the major reviewed themes, HDBSCAN outliers or general posts, and the four rolling years. Two trained coders independently identify the dominant concern in each post without being shown the model topic or candidate theme. Approximately 30 to 50 posts per major theme, plus outlier and general-post strata, is a practical starting range; the final sample size should be set prospectively based on feasible coding capacity.", size=8.9, after=3)
    paragraph(doc, "2. Freeze a human-readable codebook. Use the first sample to document missing concepts, merged concepts, inclusion and exclusion rules, and examples. Freeze the nine-theme codebook before drawing a fresh, independent validation sample. Validate the nine top-level themes before attempting the more granular 27-cluster subtheme map.", size=8.9, after=3)
    paragraph(doc, "3. Independently validate the model map. On the fresh sample, coders apply the frozen codebook while still blinded to BERTopic assignments. Report raw agreement and Krippendorff's alpha before adjudication. Then compare adjudicated human codes with the BERTopic map using per-theme precision, recall, F1, confusion matrices, and coverage of outlier or other posts.", size=8.9, after=3)
    paragraph(doc, "4. Use preset fit ratings only as a secondary audit. After independent coding is complete, representative posts and random nonmatching controls may be rated for how well a displayed model label fits. This can assess coherence and interpretability, but it should not replace blind coding because showing the proposed theme first can anchor raters.", size=8.9, after=5)
    heading(doc, "Manuscript revision status", 2)
    callout_pair(doc, "Sections revised", "The revised manuscript includes the title and abstract; current-study framing and research question; Methods, Results, Discussion, implications, limitations; figures; primary BERTopic references; and the collection-to-filtering count of 577,190 retrieved archive records followed by 97,254 analyzed posts.", "Before circulation", "Add the completed polarity method and its bounded exploratory results to the manuscript, then harmonize retained Background language with the paper's descriptive scope. Historical theory and prior turnover findings should remain context, not claims tested by this study.")
    heading(doc, "Sampling-frame framing", 2)
    callout_pair(doc, "Intended analytic focus", "The management-term filter intentionally concentrates a problem-oriented, likely negatively skewed subset of r/antiwork discussion. The purpose was to surface recurring management concerns that may warrant follow-up.", "Boundary for the manuscript", "Do not present the results as average employee sentiment or a standalone priority ranking. Frame the themes as candidate areas to examine alongside representative evidence and organizational context.")
    heading(doc, "Materials for drafting", 2)
    paragraph(doc, "Draft in this order: collection and sampling frame; coverage, theme-prevalence, and model-based polarity results; limitations; then an exploratory discussion. Use the documented theme map and avoid causal, full-conversation, or human-validated sentiment claims.", size=9.4, after=4)
    p = paragraph(doc, "Reproducible code and shareable evidence: ", size=9.4, after=4)
    hyperlink(p, REPO_URL, REPO_URL)
    paragraph(doc, "Method records: planning/SCRAPING.md, planning/PREPROCESSING.md, planning/WRITING_HANDOFF.md, planning/ROBUSTNESS.md, and planning/CANDIDATE_THEME_CODEBOOK.csv.", size=8.4, color=MUTED, after=3)
    paragraph(doc, "Current status: the topic-modeling and full-corpus polarity analyses are complete. The manuscript is revised as a descriptive, exploratory topic-prevalence study and needs the completed polarity method and results incorporated before circulation.", size=9.4, color=NAVY, bold=True)

    heading(doc, "Appendix: Reviewed cluster map")
    paragraph(doc, "The table below makes the review trail visible. It lists all 30 largest clusters examined for the candidate-theme map, their full-corpus size (n), representative terms, assigned candidate theme and subtheme, and inclusion decision.", size=9.4, after=6)
    reviewed_clusters = load_reviewed_clusters()
    for page, start in enumerate(range(0, len(reviewed_clusters), 15), start=1):
        if page > 1:
            doc.add_page_break()
        heading(doc, f"Reviewed clusters {start + 1}-{min(start + 15, len(reviewed_clusters))}", 2)
        reviewed_cluster_table(doc, reviewed_clusters[start:start + 15])
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
